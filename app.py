import builtins
import io
import json
import os
import re
import time
import uuid
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import plotly.express as px
from google import genai
from google.genai import types

# البرنامج بيجرب الموديلات دي بالترتيب لو واحد منها زحمة أو مش موجود
MODELS = ["gemini-3.8-flash", "gemini-3.5-flash", "gemini-3-flash-preview"]

# على السيرفر بنحط APP_MODE = "server" في الـ Secrets، فبيشتغل بوضع الاستضافة
SERVER_MODE = os.environ.get("APP_MODE", "").lower() == "server"

# ---------- النصوص (عربي / إنجليزي) ----------
STR = {
    "ar": {
        "title": "KAI For Data Analyst",
        "subtitle": "ارفع بياناتك، واطرح أسئلتك، واحصل على الرؤى",
        "tagline": "حوّل بياناتك إلى رؤى",
        "nav_overview": "نظرة عامة",
        "nav_upload": "رفع البيانات",
        "nav_preview": "معاينة البيانات",
        "nav_chat": "الدردشة الذكية",
        "nav_charts": "الرسوم البيانية",
        "nav_report": "تقرير التحليل",
        "up_title": "ارفع ملف Excel أو CSV",
        "up_sub": "اسحب الملف وأفلته هنا، أو انقر للاستعراض",
        "no_file_title": "لم يُرفع أي ملف بعد",
        "no_file_sub": "ارفع ملفًا لتبدأ بتحليل بياناتك.",
        "current_file": "الملف الحالي",
        "remove_file": "إزالة الملف",
        "go_overview": "عرض النظرة العامة",
        "kpi_numeric": "الأعمدة الرقمية",
        "chart1_title": "{y} حسب {x}",
        "chart2_title": "أعلى الفئات في {x}",
        "preview_title": "معاينة البيانات",
        "view_all": "عرض الكل",
        "tables_found": "الجداول المكتشفة: {n}",
        "api_label": "مفتاح Gemini (API Key)",
        "api_caption": "يبقى المفتاح على جهازك ولا يُحفظ في أي مكان.",
        "upload": "ارفع ملف Excel أو CSV",
        "no_tables": "لم أعثر على أي جدول في هذا الملف.",
        "uploaded_ok": "تم رفع الملف بنجاح! عُثر على {n} جدول.",
        "choose_table": "اختر الجدول الذي تريد العمل عليه",
        "tab_preview": "معاينة البيانات",
        "tab_analysis": "التحليل",
        "tab_charts": "الرسوم البيانية",
        "tab_summary": "الملخص الذكي",
        "tab_chat": "الدردشة",
        "rows": "عدد الصفوف",
        "cols": "عدد الأعمدة",
        "missing": "القيم المفقودة",
        "missing_per_col": "القيم المفقودة في كل عمود",
        "missing_count": "عدد القيم المفقودة",
        "num_stats": "إحصاءات الأعمدة الرقمية",
        "no_numeric": "لا توجد أعمدة رقمية في هذا الملف.",
        "need_numeric": "يلزم وجود عمود رقمي واحد على الأقل لرسم المخطط.",
        "chart_type": "نوع الرسم",
        "chart_bar": "أعمدة",
        "chart_line": "خطي",
        "chart_pie": "دائري",
        "x_col": "العمود الأفقي (X)",
        "y_col": "العمود الرقمي (Y)",
        "summary_title": "ملخص: {name}",
        "fold_open": "فتح الملخص",
        "fold_close": "طيّ الملخص",
        "del_summary": "حذف هذا الملخص",
        "need_key": "أدخل مفتاح API في الشريط الجانبي أولًا.",
        "need_key_short": "أدخل مفتاح API في الشريط الجانبي أولًا.",
        "btn_rewrite": "إعادة كتابة الملخص",
        "btn_write": "اكتب الملخص",
        "spin_summary": "يقوم Gemini بتحليل البيانات...",
        "doc_sub": "من ملف: {file}",
        "save_word": "📄 حفظ بصيغة Word",
        "save_pdf": "📕 حفظ بصيغة PDF",
        "install_libs": "لحفظ ملفات Word وPDF، ثبّت المكتبات بالأمر التالي في الشاشة السوداء:",
        "details": "التفاصيل",
        "saved_ok": "تم حفظ الملف في مجلد التنزيلات (Downloads) باسم: {name}",
        "others_exp": "ملخصات محفوظة لجداول أخرى ({n})",
        "other_from": "**{table}** (من ملف {file})",
        "chat_title": "اسأل ملفك أي سؤال",
        "chat_hint": "مثال: ما أعلى شهر في المبيعات؟ أو: ما متوسط المبيعات لكل منتج؟",
        "clear_chat": "مسح المحادثة",
        "chat_input": "اكتب سؤالك أو سجّل رسالة صوتية...",
        "spin_listen": "جارٍ الاستماع إلى الرسالة الصوتية...",
        "spin_think": "جارٍ التفكير والحساب...",
        "calc_fail": "تعذّر عليّ حساب هذه الإجابة. حاول صياغة السؤال بطريقة أخرى.",
        "result_label": "النتيجة:",
        "copy": "نسخ",
        "delete_msg": "حذف الرسالة",
        "none": "لا يوجد",
        "quota_err": (
            "لقد استهلكتَ الحصة المجانية اليومية من Gemini لجميع النماذج. "
            "حاول مجددًا بعد تجدد الحصة (غالبًا خلال بضع ساعات)، "
            "أو استخدم مفتاحًا من مشروع جديد."
        ),
        "generic_err": "حدثت مشكلة: {msg}",
        "facts": (
            "\nعدد الصفوف: {n}\nأسماء الأعمدة وأنواعها:\n{dtypes}\n\n"
            "القيم المفقودة:\n{missing}\n\nإحصاءات الأعمدة الرقمية:\n{stats}\n\n"
            "أول 15 صفًا من البيانات:\n{head}\n"
        ),
        "summary_prompt": (
            "أنت محلل بيانات خبير. فيما يلي معلومات عن ملف بيانات رفعه المستخدم:\n\n{facts}\n\n"
            "اكتب تقريرًا موجزًا وواضحًا باللغة العربية الفصحى، وخاطب القارئ بصيغة المذكر، "
            "على أن يتضمن:\n"
            "1. نظرة عامة: عمّا يتحدث الملف\n"
            "2. أهم النتائج والاتجاهات (اذكر أرقامًا حقيقية من البيانات فقط)\n"
            "3. أي مشكلات في جودة البيانات\n"
            "4. توصيات عملية\n\n"
            "لا تخترع أي رقم غير موجود في المعلومات أعلاه."
        ),
        "code_prompt": (
            "لديك جدول بيانات في pandas اسمه df.\nالأعمدة وأنواعها:\n{dtypes}\n\n"
            "أول 5 صفوف:\n{head}\n\nسؤال المستخدم: {question}\n\n"
            "اكتب كود Python فقط (دون شرح ودون import) يحسب الإجابة\n"
            "ويضع الناتج في متغير اسمه result.\n"
            "استخدم df و pd فقط، ولا تعدّل df الأصلي. اكتب الكود فقط."
        ),
        "fix_code": "\n\nالكود الذي كتبته أظهر خطأً: {error}\nصحّحه واكتب الكود فقط.",
        "answer_prompt": (
            "سؤال المستخدم: {question}\n\nنتيجة الحساب من البيانات:\n{result}\n\n"
            "اكتب إجابة قصيرة وواضحة باللغة العربية الفصحى، وخاطب القارئ بصيغة المذكر، "
            "بناءً على هذه النتيجة فقط.\n"
            "لا تخترع أي رقم غير موجود فيها."
        ),
    },
    "en": {
        "title": "KAI For Data Analyst",
        "subtitle": "Upload your data, ask questions, get insights",
        "tagline": "Turn your data into insights",
        "nav_overview": "Overview",
        "nav_upload": "Upload Dataset",
        "nav_preview": "Data Preview",
        "nav_chat": "AI Chat",
        "nav_charts": "Charts",
        "nav_report": "Analysis Report",
        "up_title": "Upload Excel or CSV",
        "up_sub": "Drag and drop your file here, or click to browse",
        "no_file_title": "No file uploaded yet",
        "no_file_sub": "Upload a file to start analyzing your data.",
        "current_file": "Current file",
        "remove_file": "Remove file",
        "go_overview": "Go to Overview",
        "kpi_numeric": "Numeric columns",
        "chart1_title": "{y} by {x}",
        "chart2_title": "Top categories in {x}",
        "preview_title": "Data Preview",
        "view_all": "View all",
        "tables_found": "Tables found: {n}",
        "api_label": "Gemini API Key",
        "api_caption": "The key stays on your computer and is not saved anywhere.",
        "upload": "Upload an Excel or CSV file",
        "no_tables": "No tables were found in this file.",
        "uploaded_ok": "File uploaded! Found {n} table(s).",
        "choose_table": "Choose the table to work on",
        "tab_preview": "Data Preview",
        "tab_analysis": "Analysis",
        "tab_charts": "Charts",
        "tab_summary": "AI Summary",
        "tab_chat": "Chat",
        "rows": "Rows",
        "cols": "Columns",
        "missing": "Missing values",
        "missing_per_col": "Missing values per column",
        "missing_count": "Missing count",
        "num_stats": "Numeric column statistics",
        "no_numeric": "There are no numeric columns in this file.",
        "need_numeric": "At least one numeric column is needed to draw a chart.",
        "chart_type": "Chart type",
        "chart_bar": "Bar",
        "chart_line": "Line",
        "chart_pie": "Pie",
        "x_col": "Horizontal column (X)",
        "y_col": "Numeric column (Y)",
        "summary_title": "Summary: {name}",
        "fold_open": "Expand summary",
        "fold_close": "Collapse summary",
        "del_summary": "Delete this summary",
        "need_key": "Enter your API Key in the sidebar (on the left) first.",
        "need_key_short": "Enter your API Key in the sidebar first.",
        "btn_rewrite": "Rewrite summary",
        "btn_write": "Write the summary",
        "spin_summary": "Gemini is analyzing the data...",
        "doc_sub": "From file: {file}",
        "save_word": "📄 Save as Word",
        "save_pdf": "📕 Save as PDF",
        "install_libs": "To save Word/PDF, install the libraries with this command in the terminal:",
        "details": "Details",
        "saved_ok": "File saved in your Downloads folder as: {name}",
        "others_exp": "Saved summaries for other tables ({n})",
        "other_from": "**{table}** (from file {file})",
        "chat_title": "Ask your file anything",
        "chat_hint": "Example: Which month has the highest sales? Or: average sales per product?",
        "clear_chat": "Clear chat",
        "chat_input": "Type your question or record a voice message...",
        "spin_listen": "Listening to the voice message...",
        "spin_think": "Thinking and calculating...",
        "calc_fail": "I couldn't calculate that answer. Try rephrasing the question.",
        "result_label": "Result:",
        "copy": "Copy",
        "delete_msg": "Delete message",
        "none": "None",
        "quota_err": (
            "You've used up the free daily quota of Gemini for all models. "
            "Try again after the quota resets (usually within a few hours), "
            "or use a key from a new project."
        ),
        "generic_err": "Something went wrong: {msg}",
        "facts": (
            "\nNumber of rows: {n}\nColumn names and types:\n{dtypes}\n\n"
            "Missing values:\n{missing}\n\nStatistics of numeric columns:\n{stats}\n\n"
            "First 15 rows of the data:\n{head}\n"
        ),
        "summary_prompt": (
            "You are an expert data analyst. Here is information about a data file "
            "the user uploaded:\n\n{facts}\n\n"
            "Write a short, clear report in English with:\n"
            "1. Overview: what the file is about\n"
            "2. Key findings and trends (cite only real numbers from the data)\n"
            "3. Any data quality issues\n"
            "4. Practical recommendations\n\n"
            "Do not invent any number that is not in the information above."
        ),
        "code_prompt": (
            "You have a pandas DataFrame named df.\nColumns and types:\n{dtypes}\n\n"
            "First 5 rows:\n{head}\n\nUser question: {question}\n\n"
            "Write Python code only (no explanation, no imports) that computes the answer\n"
            "and puts the result in a variable named result.\n"
            "Use only df and pd, and do not modify the original df. Output the code only."
        ),
        "fix_code": "\n\nThe code you wrote raised an error: {error}\nFix it and output the code only.",
        "answer_prompt": (
            "User question: {question}\n\nCalculation result from the data:\n{result}\n\n"
            "Write a short, clear answer in English based only on this result.\n"
            "Do not invent any number that is not in it."
        ),
    },
}


def t(key, **kw):
    """بيرجّع النص باللغة الحالية (عربي أو إنجليزي)."""
    lang = st.session_state.get("lang", "ar")
    text = STR[lang].get(key) or STR["ar"].get(key, key)
    return text.format(**kw) if kw else text


def disp(name):
    """اسم الجدول للعرض (جدول -> Table في الإنجليزي)."""
    if st.session_state.get("lang") == "en":
        return str(name).replace(" - جدول ", " - Table ")
    return str(name)


def ui_rtl():
    return st.session_state.get("lang", "ar") == "ar"


def is_rtl_text(text):
    return bool(re.search(r"[\u0600-\u06FF]", text))


st.set_page_config(page_title="KAI For Data Analyst", layout="wide")
st.session_state.setdefault("lang", "ar")

RTL_CSS = """
<style>
[data-testid="stMainBlockContainer"],
[data-testid="stSidebarContent"] { direction: rtl; }

[data-testid="stMainBlockContainer"] :is(h1, h2, h3, h4, h5, h6, p, li, label, summary, span, div),
[data-testid="stSidebarContent"] :is(h1, h2, h3, h4, h5, h6, p, li, label, summary, span, div) {
    text-align: right !important;
}
[data-testid="stMainBlockContainer"] :is(h1, h2, h3, h4, h5, h6, p, li, label, summary),
[data-testid="stSidebarContent"] :is(h1, h2, h3, h4, h5, h6, p, li, label, summary) {
    direction: rtl !important;
}

/* قوائم الاختيار (بتتفتح برا الصفحة) */
[data-baseweb="popover"], [data-baseweb="popover"] *,
[role="listbox"], [role="listbox"] *, [role="option"], [role="option"] *,
[data-testid="stSelectboxVirtualDropdown"], [data-testid="stSelectboxVirtualDropdown"] * {
    direction: rtl !important;
    text-align: right !important;
}

/* حاجات لازم تفضل من الشمال */
[data-testid="stDataFrame"], [data-testid="stPlotlyChart"],
[data-testid="stSidebar"] input[type="password"] { direction: ltr; }
[data-testid="stChatInput"] textarea { direction: rtl; text-align: right; }
</style>
"""
BASE_CSS = """
<style>
:root { --navy:#132440; --navy2:#183052; --teal:#168982; --teal-d:#0f807c; --teal-l:#4fc4b5;
        --mint:#c6e9e6; --bg:#f7f9fc; --line:#e6ebf2; --muted:#6b7a90; }
.stApp { background: var(--bg); }
[data-testid="stHeader"] { background: transparent; }
[data-testid="stAppDeployButton"] { display: none; }
[data-testid="stMainBlockContainer"] { padding-top: 1.6rem; max-width: 1500px; }

/* ===== الشريط الجانبي ===== */
[data-testid="stSidebar"] { background: var(--navy); }
.brand { display:flex; align-items:center; gap:.7rem; padding:.3rem .3rem 1.2rem; }
.brand-logo { width:40px; height:40px; border-radius:12px; color:#fff; font-weight:800; font-size:1.35rem;
              background:linear-gradient(135deg,var(--teal-l),var(--teal)); display:flex;
              align-items:center; justify-content:center; }
.brand-name { color:#fff; font-size:1.6rem; font-weight:700; letter-spacing:.5px; }
.side-gap { height:1.4rem; }
.side-tagline { color:#9fb3cf; font-size:.85rem; margin-top:1rem; }

[data-testid="stSidebar"] [role="radiogroup"] { gap:.3rem; }
[data-testid="stSidebar"] label[data-testid="stRadioOption"],
[data-testid="stSidebar"] label[data-baseweb="radio"] { width:100%; padding:.75rem 1rem; margin:0;
        border-radius:12px; cursor:pointer; transition:background .15s; }
[data-testid="stSidebar"] label[data-testid="stRadioOption"]:hover,
[data-testid="stSidebar"] label[data-baseweb="radio"]:hover { background:rgba(255,255,255,.07); }
/* إخفاء دايرة الراديو */
[data-testid="stSidebar"] label[data-testid="stRadioOption"] > div > div:first-child,
[data-testid="stSidebar"] label[data-baseweb="radio"] > div:first-child { display:none !important; }
[data-testid="stSidebar"] label[data-testid="stRadioOption"][data-selected="true"],
[data-testid="stSidebar"] label[data-baseweb="radio"]:has(input:checked) {
        background:linear-gradient(90deg,#1f6b77,#1b8587); }
[data-testid="stSidebar"] label[data-testid="stRadioOption"] p,
[data-testid="stSidebar"] label[data-baseweb="radio"] p { color:#dce6f5; font-size:1.05rem; }
[data-testid="stSidebar"] label[data-testid="stRadioOption"][data-selected="true"] p,
[data-testid="stSidebar"] label[data-baseweb="radio"]:has(input:checked) p { color:#fff; font-weight:700; }

/* ===== الكروت ===== */
[data-testid="stVerticalBlockBorderWrapper"], div[data-testid="stVerticalBlock"][class*="st-key-upcard"] {
        background:#fff; border:1px solid var(--line) !important; border-radius:16px;
        box-shadow:0 2px 10px rgba(19,36,64,.05); }
.kpi { display:flex; align-items:center; gap:1rem; background:#fff; border:1px solid var(--line);
       border-radius:16px; padding:1.1rem 1.2rem; box-shadow:0 2px 10px rgba(19,36,64,.05);
       margin-bottom:1rem; }
.kpi-icon { width:52px; height:52px; border-radius:50%; display:flex; align-items:center;
            justify-content:center; flex-shrink:0; }
.kpi-icon.teal { background:#d7f1ee; color:#168982; }
.kpi-icon.blue { background:#dfeafb; color:#2f6fd1; }
.kpi-icon.amber { background:#fdeccd; color:#c27a0a; }
.kpi-icon.violet { background:#ebe3fb; color:#6a4bc4; }
.kpi-label { color:var(--navy2); font-size:.95rem; font-weight:500; }
.kpi-value { color:var(--navy); font-size:1.9rem; font-weight:800; line-height:1.15; }
.page-title { font-size:2.1rem; font-weight:800; color:var(--navy2); line-height:1.2; }
.page-sub { color:var(--muted); font-size:1rem; margin-top:.15rem; margin-bottom:1.1rem; }
.card-title { font-size:1.15rem; font-weight:700; color:var(--navy2); margin-bottom:.5rem; }
.file-chip { display:inline-block; background:#fff; border:1px solid var(--line); border-radius:12px;
             padding:.55rem .9rem; color:var(--navy2); font-weight:600; }
.empty-title { font-size:1.2rem; font-weight:700; color:var(--navy2); }
.empty-sub { color:var(--muted); padding-bottom:.9rem; }
.up-icon { font-size:2.4rem; color:var(--teal); text-align:center; line-height:1; }
.up-title { font-weight:700; text-align:center; color:var(--navy2); font-size:1.15rem; margin-top:.3rem; }
.up-sub { text-align:center; color:var(--muted); margin-bottom:.8rem; }
[data-testid="stMainBlockContainer"] .up-icon, [data-testid="stMainBlockContainer"] .up-title,
[data-testid="stMainBlockContainer"] .up-sub { text-align:center !important; }

/* ===== أزرار وحقول ===== */
.stButton > button, .stDownloadButton > button { border-radius:10px; border:1px solid var(--line);
        background:#fff; color:var(--navy2); font-weight:600; }
.stButton > button:hover, .stDownloadButton > button:hover { border-color:var(--teal); color:var(--teal); }
[data-testid="stFileUploaderDropzone"] { border:2px dashed var(--mint); background:#f3fbfa; border-radius:14px; }
[data-testid="stFileUploaderDropzone"] button { background:var(--teal); color:#fff; border:none; }
[data-testid="stSidebar"] input { background:#1c3354; color:#fff; }

/* ===== الدردشة ===== */
[data-testid="stChatMessage"] { background:#edf2fa; border-radius:14px; padding:.8rem 1rem; }
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) { background:#e6f2fd; }
[data-testid="stChatInput"] { border-radius:14px; }
</style>
"""

st.markdown(BASE_CSS, unsafe_allow_html=True)
if st.session_state.lang == "ar":
    st.markdown(RTL_CSS, unsafe_allow_html=True)


def toggle_lang():
    st.session_state.lang = "en" if st.session_state.lang == "ar" else "ar"


# ---------- الشريط الجانبي: الشعار + التنقل + مفتاح Gemini ----------
NAV = [
    ("overview", "home"), ("upload", "cloud_upload"), ("preview", "table_chart"),
    ("chat", "chat"), ("charts", "bar_chart"), ("report", "description"),
]
NAV_ICONS = dict(NAV)

st.sidebar.markdown(
    f"""
    <div class="brand">
      <div class="brand-logo">K</div>
      <div class="brand-name">KAI</div>
    </div>
    """,
    unsafe_allow_html=True,
)
page = st.sidebar.radio(
    "nav",
    [k for k, _ in NAV],
    format_func=lambda k: f":material/{NAV_ICONS[k]}:  {t('nav_' + k)}",
    key="nav",
    label_visibility="collapsed",
)
st.sidebar.markdown("<div class='side-gap'></div>", unsafe_allow_html=True)
api_key = st.sidebar.text_input(
    t("api_label"), type="password", value=st.session_state.get("api_key_saved", "")
)
st.session_state["api_key_saved"] = api_key
st.sidebar.caption(t("api_caption"))
st.sidebar.markdown(
    f"<div class='side-tagline'>✦ {t('tagline')}</div>", unsafe_allow_html=True
)


# ---------- دوال مساعدة للدردشة ----------
def ask_gemini(client, prompt):
    """بيجرب الموديلات بالترتيب، وبيعيد المحاولة لو الموديل زحمة (503)."""
    last_error = None
    for model in MODELS:
        for attempt in range(3):
            try:
                response = client.models.generate_content(model=model, contents=prompt)
                return response.text
            except Exception as e:
                last_error = e
                msg = str(e)
                if "503" in msg or "UNAVAILABLE" in msg:
                    time.sleep(3 * (attempt + 1))  # استنى وجرب تاني
                    continue
                if (
                    "404" in msg or "NOT_FOUND" in msg
                    or "429" in msg or "RESOURCE_EXHAUSTED" in msg
                ):
                    break  # الموديل مش متاح أو حصته خلصت، جرب اللي بعده
                raise
    raise last_error


def friendly_error(e):
    msg = str(e)
    if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
        return t("quota_err")
    return t("generic_err", msg=msg)


def transcribe(client, audio_file):
    """بيحوّل الرسالة الصوتية لنص باستخدام Gemini."""
    audio_bytes = audio_file.getvalue() if hasattr(audio_file, "getvalue") else audio_file.read()
    part = types.Part.from_bytes(data=audio_bytes, mime_type="audio/wav")
    text = ask_gemini(
        client,
        [
            "Transcribe the speech in this audio exactly as spoken (usually Arabic or English). "
            "Output only the text, with no explanation.",
            part,
        ],
    )
    return text.strip()


def clean_code(text):
    """بيشيل علامات ``` لو Gemini حطها حوالين الكود."""
    text = text.strip()
    text = re.sub(r"^```(?:python)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    return text


FORBIDDEN = [
    r"\bimport\b", r"\bopen\s*\(", r"\bexec\b", r"\beval\b", r"__",
    r"\bos\b", r"\bsys\b", r"\bsubprocess\b", r"\binput\s*\(",
    r"\bcompile\b", r"\bglobals\b", r"\blocals\b", r"\bgetattr\b",
    r"\bsetattr\b", r"\bdelattr\b", r"\bread_\w+", r"\bto_\w+\s*\(",
]
ALLOWED_BUILTINS = [
    "len", "sum", "min", "max", "abs", "round", "sorted", "range", "list",
    "dict", "set", "tuple", "str", "int", "float", "bool", "enumerate",
    "zip", "map", "filter", "any", "all",
]


def run_code(code, df):
    """بينفذ الكود في بيئة مقيّدة (من غير import ولا ملفات)."""
    for pattern in FORBIDDEN:
        if re.search(pattern, code):
            raise ValueError("يحتوي الكود على أمر غير مسموح به لأسباب أمنية.")
    safe_builtins = {n: getattr(builtins, n) for n in ALLOWED_BUILTINS}
    env = {"df": df.copy(), "pd": pd, "__builtins__": safe_builtins}
    exec(code, env)
    if "result" not in env:
        raise ValueError("لم يضع الكود الناتج في متغير result.")
    return env["result"]


def result_to_text(result):
    if isinstance(result, (pd.DataFrame, pd.Series)):
        text = result.to_string()
    else:
        text = str(result)
    return text[:3000]


# ---------- دوال تقسيم الشيت لجداول ----------
def _split_axis(frame, axis):
    """بيقسم على الصفوف الفاضية (axis=0) أو الأعمدة الفاضية (axis=1)."""
    empty = frame.isna().all(axis=1) if axis == 0 else frame.isna().all(axis=0)
    parts, start = [], None
    for i, is_empty in enumerate(empty.tolist()):
        if not is_empty and start is None:
            start = i
        elif is_empty and start is not None:
            parts.append((start, i))
            start = None
    if start is not None:
        parts.append((start, len(empty)))
    if axis == 0:
        return [frame.iloc[a:b, :] for a, b in parts]
    return [frame.iloc[:, a:b] for a, b in parts]


def split_blocks(frame):
    """بيلاقي كل الجداول اللي في الشيت (بينهم صف أو عمود فاضي)."""
    # نشيل الصفوف والأعمدة الفاضية من الأطراف بس (مش اللي في النص)
    rows = frame.notna().any(axis=1).to_numpy().nonzero()[0]
    cols = frame.notna().any(axis=0).to_numpy().nonzero()[0]
    if len(rows) == 0 or len(cols) == 0:
        return []
    frame = frame.iloc[rows[0]:rows[-1] + 1, cols[0]:cols[-1] + 1]
    for axis in (0, 1):
        parts = _split_axis(frame, axis)
        if len(parts) > 1:
            result = []
            for part in parts:
                result += split_blocks(part)
            return result
    return [frame]


def block_to_df(block):
    """بيحوّل الجزء ده لجدول: أول صف عناوين، والباقي بيانات."""
    block = block.reset_index(drop=True)
    header = [
        str(h) if pd.notna(h) else f"عمود {i + 1}"
        for i, h in enumerate(block.iloc[0].tolist())
    ]
    seen, names = {}, []
    for h in header:  # لو في عناوين مكررة
        if h in seen:
            seen[h] += 1
            names.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 0
            names.append(h)
    out = block.iloc[1:].reset_index(drop=True)
    out.columns = names
    return out.infer_objects()


def load_tables(uploaded):
    """بيرجّع قاموس: اسم الجدول -> الجدول (من كل الشيتات)."""
    if uploaded.name.endswith(".csv"):
        sheets = {"CSV": pd.read_csv(uploaded, header=None, skip_blank_lines=False)}
    else:
        sheets = pd.read_excel(uploaded, sheet_name=None, header=None)

    tables = {}
    for sheet_name, raw in sheets.items():
        blocks = [b for b in split_blocks(raw) if len(b) >= 2]
        for i, block in enumerate(blocks, start=1):
            label = f"{sheet_name} - جدول {i}" if len(blocks) > 1 else str(sheet_name)
            tables[label] = block_to_df(block)
    return tables


# ---------- حفظ الملخصات على الجهاز ----------
SUMMARIES_FILE = Path(__file__).parent / "summaries.json"


def load_summaries():
    if SERVER_MODE:
        return {}
    try:
        return json.loads(SUMMARIES_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_summaries(data):
    if SERVER_MODE:
        return
    SUMMARIES_FILE.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def delete_summary(key):
    st.session_state.summaries.pop(key, None)
    st.session_state.setdefault("collapsed_summaries", set()).discard(key)
    st.session_state.setdefault("open_saved", set()).discard(key)
    save_summaries(st.session_state.summaries)


def toggle_saved(key):
    """فتح/قفل ملخص محفوظ في قايمة الملخصات التانية."""
    st.session_state.setdefault("open_saved", set()).symmetric_difference_update({key})


def wide_button(label, **kwargs):
    try:
        return st.button(label, width="stretch", **kwargs)
    except TypeError:  # نسخة Streamlit قديمة
        return st.button(label, use_container_width=True, **kwargs)


def safe_download_button(container, label, **kwargs):
    """زرار تحميل مبيعملش إعادة تشغيل للصفحة."""
    try:
        return container.download_button(label, on_click="ignore", **kwargs)
    except TypeError:
        return container.download_button(label, **kwargs)


def toggle_summary(key):
    """طي/فتح الملخص."""
    st.session_state.setdefault("collapsed_summaries", set()).symmetric_difference_update({key})


def save_to_downloads(data, table_name, ext):
    """بيحفظ الملف في فولدر Downloads باسم الجدول."""
    folder = Path.home() / "Downloads"
    folder.mkdir(exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", str(table_name)).strip(" .") or "ملخص"
    path = folder / f"{safe}.{ext}"
    n = 2
    while path.exists():  # لو الاسم موجود، نضيف رقم بدل ما نمسح القديم
        path = folder / f"{safe} ({n}).{ext}"
        n += 1
    path.write_bytes(data)
    return path


# ---------- تصدير الملخص (Word / PDF) ----------
def parse_markdown(text):
    """بيحوّل نص الملخص لقائمة (نوع، نص): عناوين ونقط وفقرات."""
    blocks = []
    for raw in text.splitlines():
        line = raw.strip().replace("`", "")
        if not line:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            blocks.append((f"h{min(len(m.group(1)), 3)}", m.group(2)))
            continue
        m = re.match(r"^[-*•]\s+(.*)", line)
        if m:
            blocks.append(("bullet", "• " + m.group(1)))
            continue
        blocks.append(("p", line))
    return blocks


def summary_to_docx(title, subtitle, text, rtl=True):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()

    def add_par(content, size=12, bold_all=False):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        if rtl:
            pPr.append(OxmlElement("w:bidi"))  # اتجاه الفقرة من اليمين لليسار
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:after"), "120")
        pPr.append(sp)
        for part in re.split(r"(\*\*.+?\*\*)", content):
            if not part:
                continue
            bold = bold_all or (part.startswith("**") and part.endswith("**"))
            run = p.add_run(part.strip("*") if part.startswith("**") else part)
            rPr = run._r.get_or_add_rPr()
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), "Arial")
            rPr.append(rf)
            if bold:
                rPr.append(OxmlElement("w:b"))
                rPr.append(OxmlElement("w:bCs"))
            for tag in ("w:sz", "w:szCs"):
                el = OxmlElement(tag)
                el.set(qn("w:val"), str(int(size * 2)))
                rPr.append(el)
            if rtl:
                rPr.append(OxmlElement("w:rtl"))

    add_par(title, size=20, bold_all=True)
    add_par(subtitle, size=10)
    sizes = {"h1": 18, "h2": 16, "h3": 14}
    for kind, content in parse_markdown(text):
        if kind in sizes:
            add_par(content, size=sizes[kind], bold_all=True)
        else:
            add_par(content, size=12)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _find_arabic_font():
    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\tahoma.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        str(Path(__file__).parent / "fonts" / "DejaVuSans.ttf"),
    ]
    for path in candidates:
        if Path(path).exists():
            return path
    raise FileNotFoundError("لم أعثر على خط يدعم اللغة العربية على هذا الجهاز.")


def summary_to_pdf(title, subtitle, text, rtl=True):
    import arabic_reshaper
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    try:
        from bidi import get_display
    except ImportError:
        from bidi.algorithm import get_display

    font = "ArabicFont"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, _find_arabic_font()))

    # بعض نسخ مكتبة bidi مبتقلبش شكل الأقواس في النص العربي، فبنتأكد ونصلّحها
    swap = str.maketrans("()[]{}<>", ")(][}{><")
    needs_swap = get_display("(ب)", base_dir="R") != "(ب)"

    def vis(t):  # تشكيل الحروف العربي + ترتيب الاتجاه
        if not rtl:
            return t
        out = get_display(arabic_reshaper.reshape(t), base_dir="R")
        return out.translate(swap) if needs_swap else out

    width, height = A4
    margin = 50
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    state = {"y": height - margin}

    def draw(content, size=12, indent=0, gap=6):
        content = content.replace("**", "")
        line_h = size * 1.6
        max_w = width - 2 * margin - indent
        lines, cur = [], ""
        for w in content.split():
            trial = (cur + " " + w).strip()
            if pdfmetrics.stringWidth(vis(trial), font, size) <= max_w:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        for ln in lines:
            if state["y"] - line_h < margin:
                c.showPage()
                state["y"] = height - margin
            c.setFont(font, size)
            if rtl:
                c.drawRightString(width - margin - indent, state["y"] - size, vis(ln))
            else:
                c.drawString(margin + indent, state["y"] - size, ln)
            state["y"] -= line_h
        state["y"] -= gap

    draw(title, size=20, gap=2)
    draw(subtitle, size=10, gap=14)
    sizes = {"h1": 18, "h2": 16, "h3": 14}
    for kind, content in parse_markdown(text):
        if kind in sizes:
            draw(content, size=sizes[kind], gap=8)
        elif kind == "bullet":
            draw(content, size=12, indent=12, gap=4)
        else:
            draw(content, size=12)
    c.save()
    return buf.getvalue()


# ---------- أدوات رسائل الشات (نسخ + حذف) ----------
CHAT_CSS = """
<style>
[class*="st-key-actions_"] { opacity: 0; transition: opacity .15s ease-in-out; }
[data-testid="stChatMessage"]:hover [class*="st-key-actions_"],
[class*="st-key-actions_"]:focus-within { opacity: 1; }
[class*="st-key-actions_"] button { padding: 0.1rem 0.6rem; min-height: 2rem; }
</style>
"""


def copy_button(text):
    """زرار نسخ بيشتغل في المتصفح مباشرة (من غير سكرول)."""
    safe = json.dumps(text).replace("</", "<\\/")
    html = f"""
    <style>
      html, body {{ margin: 0; padding: 0; overflow: hidden; background: transparent; }}
      button {{ background: transparent; border: 1px solid #666; border-radius: 8px;
                cursor: pointer; font-size: 15px; height: 34px; padding: 0 12px;
                display: block; margin: 2px 0 0 auto; }}
    </style>
    <button id="b" title="{t('copy')}">📋</button>
    <script>
    const t = {safe};
    const b = document.getElementById("b");
    b.onclick = async () => {{
        try {{ await navigator.clipboard.writeText(t); }}
        catch (e) {{
            const a = document.createElement("textarea");
            a.value = t; document.body.appendChild(a); a.select();
            document.execCommand("copy"); a.remove();
        }}
        b.textContent = "✅";
        setTimeout(() => {{ b.textContent = "📋"; }}, 1200);
    }};
    </script>
    """
    if hasattr(st, "iframe"):  # النسخ الجديدة من Streamlit
        st.iframe(html, height=42)
    else:  # النسخ الأقدم
        components.html(html, height=42, scrolling=False)


def delete_message(msg_id):
    st.session_state.messages = [
        m for m in st.session_state.messages if m.get("id") != msg_id
    ]


def render_message(m):
    m.setdefault("id", uuid.uuid4().hex)
    with st.chat_message(m["role"]):
        st.markdown(m["content"])
        with st.container(key=f"actions_{m['id']}"):
            if ui_rtl():
                c1, c2, _ = st.columns([1, 1, 14], vertical_alignment="center")
            else:
                _, c1, c2 = st.columns([14, 1, 1], vertical_alignment="center")
            with c1:
                copy_button(m["content"])
            with c2:
                st.button(
                    "🗑",
                    key=f"del_{m['id']}",
                    help=t("delete_msg"),
                    on_click=delete_message,
                    args=(m["id"],),
                )


# ---------- البيانات المحفوظة في الجلسة ----------
@st.cache_data(show_spinner=False)
def parse_tables(data, name):
    f = io.BytesIO(data)
    f.name = name
    return load_tables(f)


def save_upload(up):
    """بيحفظ الملف في الجلسة عشان ميتمسحش لما تغيري الصفحة أو اللغة."""
    data = up.getvalue()
    if st.session_state.get("file_bytes") == data and st.session_state.get("file_name") == up.name:
        return False
    st.session_state.update(
        file_bytes=data, file_name=up.name, current_table=None, messages=[]
    )
    return True


def remove_file():
    for k in ("file_bytes", "file_name", "current_table"):
        st.session_state.pop(k, None)
    st.session_state["messages"] = []
    st.session_state["uploader_n"] = st.session_state.get("uploader_n", 0) + 1


def go_to(page_key):
    st.session_state["nav"] = page_key


def uploader_card(suffix):
    with st.container(border=True, key=f"upcard_{suffix}"):
        st.markdown(
            f"<div class='up-icon'>☁</div><div class='up-title'>{t('up_title')}</div>"
            f"<div class='up-sub'>{t('up_sub')}</div>",
            unsafe_allow_html=True,
        )
        up = st.file_uploader(
            "📂 Excel / CSV",
            type=["csv", "xlsx"],
            key=f"uploader_{st.session_state.get('uploader_n', 0)}_{suffix}",
            label_visibility="collapsed",
        )
        if up is not None and save_upload(up):
            st.rerun()


# ---------- أدوات الشكل: كروت وأيقونات ورسوم ----------
SVG = (
    '<svg viewBox="0 0 24 24" width="22" height="22" fill="none" stroke="currentColor" '
    'stroke-width="2" stroke-linecap="round" stroke-linejoin="round">{}</svg>'
)
ICON_ROWS = SVG.format('<ellipse cx="12" cy="5" rx="8" ry="3"/><path d="M4 5v6c0 1.7 3.6 3 8 3s8-1.3 8-3V5M4 11v6c0 1.7 3.6 3 8 3s8-1.3 8-3v-6"/>')
ICON_COLS = SVG.format('<rect x="3" y="4" width="18" height="16" rx="2"/><path d="M9 4v16M15 4v16"/>')
ICON_MISS = SVG.format('<circle cx="12" cy="12" r="9"/><path d="M12 8v5M12 16.5v.5"/>')
ICON_NUM = SVG.format('<path d="M5 9h14M5 15h14M10 4L8 20M16 4l-2 16"/>')
TEAL_SEQ = ["#168982", "#4fc4b5", "#0f807c", "#8bc4bc", "#c6e9e6", "#183052"]


def kpi(col, icon, label, value, tone="teal"):
    col.markdown(
        f'<div class="kpi"><div class="kpi-icon {tone}">{icon}</div>'
        f'<div class="kpi-body"><div class="kpi-label">{label}</div>'
        f'<div class="kpi-value">{value}</div></div></div>',
        unsafe_allow_html=True,
    )


def card_title(text):
    st.markdown(f'<div class="card-title">{text}</div>', unsafe_allow_html=True)


def style_fig(fig):
    fig.update_layout(
        margin=dict(l=10, r=10, t=10, b=10),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color="#183052"),
        colorway=TEAL_SEQ,
        height=320,
    )
    fig.update_xaxes(showgrid=False)
    fig.update_yaxes(gridcolor="#e6ebf2")
    return fig


def show_chart(fig):
    try:
        st.plotly_chart(fig, width="stretch")
    except TypeError:
        st.plotly_chart(fig, use_container_width=True)


def show_df(frame):
    try:
        st.dataframe(frame, width="stretch")
    except TypeError:
        st.dataframe(frame, use_container_width=True)


# ---------- تحميل الجداول من الملف المحفوظ ----------
tables = {}
if st.session_state.get("file_bytes"):
    tables = parse_tables(st.session_state["file_bytes"], st.session_state["file_name"])
file_name = st.session_state.get("file_name")

# ---------- الهيدر: العنوان + اختيار الجدول + زرار اللغة ----------
h_title, h_table, h_lang = st.columns([5, 3, 1], vertical_alignment="center")
h_title.markdown(
    f'<div class="page-title">{t("title")}</div><div class="page-sub">{t("subtitle")}</div>',
    unsafe_allow_html=True,
)
df, table_name = None, None
if tables:
    keys = list(tables.keys())
    if len(keys) > 1:
        labels = [disp(k) for k in keys]
        saved = st.session_state.get("current_table")
        chosen = h_table.selectbox(
            t("choose_table"), labels,
            index=keys.index(saved) if saved in keys else 0,
            label_visibility="collapsed",
        )
        table_name = keys[labels.index(chosen)]
    else:
        table_name = keys[0]
        h_table.markdown(f'<div class="file-chip">📄 {file_name}</div>', unsafe_allow_html=True)
    df = tables[table_name]
    if st.session_state.get("current_table") != table_name:
        st.session_state["current_table"] = table_name
        st.session_state["messages"] = []
elif st.session_state.get("file_bytes"):
    st.error(t("no_tables"))

h_lang.button(
    "EN" if st.session_state.lang == "ar" else "AR",
    key="lang_btn",
    on_click=toggle_lang,
    help="Switch language / تغيير اللغة",
)

numeric_cols = df.select_dtypes(include="number").columns.tolist() if df is not None else []

# ---------- صفحة رفع البيانات + حالة عدم وجود ملف ----------
if df is None and page != "upload":
    with st.container(border=True):
        st.markdown(
            f'<div class="empty-title">{t("no_file_title")}</div>'
            f'<div class="empty-sub">{t("no_file_sub")}</div>',
            unsafe_allow_html=True,
        )
    uploader_card("empty")

if page == "upload":
    uploader_card("main")
    if df is not None:
        with st.container(border=True):
            card_title(t("current_file"))
            st.markdown(f'<div class="file-chip">📄 {file_name}</div>', unsafe_allow_html=True)
            st.caption(t("tables_found", n=len(tables)))
            b1, b2, _ = st.columns([1, 1, 3])
            b1.button(t("go_overview"), key="go_ov", on_click=go_to, args=("overview",))
            b2.button(t("remove_file"), key="rm_file", on_click=remove_file)

if df is not None:
    # ---------- نظرة عامة ----------
    if page == "overview":
        k1, k2, k3, k4 = st.columns(4)
        kpi(k1, ICON_ROWS, t("rows"), f"{len(df):,}", "teal")
        kpi(k2, ICON_COLS, t("cols"), len(df.columns), "blue")
        kpi(k3, ICON_MISS, t("missing"), int(df.isna().sum().sum()), "amber")
        kpi(k4, ICON_NUM, t("kpi_numeric"), len(numeric_cols), "violet")

        cat_cols = [c for c in df.columns if c not in numeric_cols]
        if numeric_cols:
            y = numeric_cols[0]
            x = cat_cols[0] if cat_cols else None
            left, right = st.columns(2)
            with left:
                with st.container(border=True):
                    card_title(t("chart1_title", y=y, x=x if x is not None else "#"))
                    fig = px.line(df, x=x, y=y) if x is not None else px.line(df, y=y)
                    fig.update_traces(line_color="#168982", fill="tozeroy", fillcolor="rgba(22,137,130,0.12)")
                    show_chart(style_fig(fig))
            with right:
                with st.container(border=True):
                    if x is not None:
                        card_title(t("chart2_title", x=x))
                        top = (
                            df.groupby(x)[y].sum().sort_values(ascending=False).head(8).reset_index()
                        )
                        fig2 = px.bar(top, x=x, y=y)
                        fig2.update_traces(marker_color="#239189")
                        show_chart(style_fig(fig2))
                    else:
                        card_title(t("num_stats"))
                        show_df(df[numeric_cols].describe())
        else:
            st.info(t("need_numeric"))

        with st.container(border=True):
            p_title, p_btn = st.columns([6, 1], vertical_alignment="center")
            p_title.markdown(f'<div class="card-title">{t("preview_title")}</div>', unsafe_allow_html=True)
            p_btn.button(t("view_all"), key="view_all", on_click=go_to, args=("preview",))
            show_df(df.head(8))

    # ---------- معاينة البيانات + التحليل ----------
    if page == "preview":
        with st.container(border=True):
            card_title(t("preview_title"))
            show_df(df.head(50))
        with st.container(border=True):
            card_title(t("missing_per_col"))
            show_df(df.isna().sum().rename(t("missing_count")))
        with st.container(border=True):
            card_title(t("num_stats"))
            if numeric_cols:
                show_df(df[numeric_cols].describe())
            else:
                st.info(t("no_numeric"))

    # ---------- الرسوم البيانية ----------
    if page == "charts":
        with st.container(border=True):
            if not numeric_cols:
                st.info(t("need_numeric"))
            else:
                chart_keys = ["bar", "line", "pie"]
                chart_labels = [t("chart_" + k) for k in chart_keys]
                c_a, c_b, c_c = st.columns(3)
                chart_type = chart_keys[chart_labels.index(c_a.selectbox(t("chart_type"), chart_labels))]
                x_col = c_b.selectbox(t("x_col"), df.columns.tolist())
                y_col = c_c.selectbox(t("y_col"), numeric_cols)

                if chart_type == "bar":
                    fig = px.bar(df, x=x_col, y=y_col)
                    fig.update_traces(marker_color="#239189")
                elif chart_type == "line":
                    fig = px.line(df, x=x_col, y=y_col)
                    fig.update_traces(line_color="#168982")
                else:
                    fig = px.pie(df, names=x_col, values=y_col)
                show_chart(style_fig(fig))

    # ---------- التبويب 4: الملخص الذكي ----------
    if page == "report":
        # الملخصات بتتحفظ في ملف summaries.json جنب app.py، فمبتتمسحش
        if "summaries" not in st.session_state:
            st.session_state.summaries = load_summaries()
        summaries = st.session_state.summaries
        for k, v in list(summaries.items()):
            if isinstance(v, str):  # تحويل الصيغة القديمة
                f_name, _, t_name = k.partition("|")
                summaries[k] = {"table": t_name.strip() or k, "file": f_name.strip(), "text": v}

        sum_key = f"{file_name} | {table_name}"
        current = summaries.get(sum_key)

        collapsed_set = st.session_state.setdefault("collapsed_summaries", set())
        is_collapsed = sum_key in collapsed_set

        title_col, fold_col, del_col = st.columns([12, 1, 1], vertical_alignment="center")
        title_col.subheader(t("summary_title", name=disp(table_name)))
        if current:
            fold_col.button(
                ":material/expand_more:" if is_collapsed else ":material/expand_less:",
                key="fold_summary",
                help=t("fold_open") if is_collapsed else t("fold_close"),
                on_click=toggle_summary, args=(sum_key,),
            )
            del_col.button(
                "🗑", key="del_summary", help=t("del_summary"),
                on_click=delete_summary, args=(sum_key,),
            )

        if not api_key:
            st.warning(t("need_key"))

        if st.button(t("btn_rewrite") if current else t("btn_write"), disabled=not api_key):
            facts = t(
                "facts",
                n=len(df),
                dtypes=df.dtypes.to_string(),
                missing=df.isna().sum().to_string(),
                stats=df[numeric_cols].describe().to_string() if numeric_cols else t("none"),
                head=df.head(15).to_string(),
            )
            prompt = t("summary_prompt", facts=facts)

            try:
                client = genai.Client(api_key=api_key)
                with st.spinner(t("spin_summary")):
                    text = ask_gemini(client, prompt)
                summaries[sum_key] = {"table": table_name, "file": file_name, "text": text}
                save_summaries(summaries)
                collapsed_set.discard(sum_key)
                st.rerun()  # نعيد الرسم عشان زراير الطي والحذف تظهر فورًا
            except Exception as e:
                st.error(friendly_error(e))

        if current and not is_collapsed:
            st.markdown(current["text"])
            doc_title = t("summary_title", name=disp(table_name))
            doc_sub = t("doc_sub", file=file_name)
            doc_rtl = is_rtl_text(current["text"])
            d1, d2, _ = st.columns([1, 1, 3])
            saved_path = None
            try:
                if SERVER_MODE:
                    safe_name = re.sub(r'[\\/:*?"<>|]', "_", str(table_name)).strip(" .") or "summary"
                    safe_download_button(
                        d1, t("save_word"),
                        data=summary_to_docx(doc_title, doc_sub, current["text"], rtl=doc_rtl),
                        file_name=f"{safe_name}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                    safe_download_button(
                        d2, t("save_pdf"),
                        data=summary_to_pdf(doc_title, doc_sub, current["text"], rtl=doc_rtl),
                        file_name=f"{safe_name}.pdf",
                        mime="application/pdf",
                    )
                else:
                    if d1.button(t("save_word"), key="save_word"):
                        saved_path = save_to_downloads(
                            summary_to_docx(doc_title, doc_sub, current["text"], rtl=doc_rtl),
                            table_name, "docx",
                        )
                    if d2.button(t("save_pdf"), key="save_pdf"):
                        saved_path = save_to_downloads(
                            summary_to_pdf(doc_title, doc_sub, current["text"], rtl=doc_rtl),
                            table_name, "pdf",
                        )
            except Exception as e:
                st.info(
                    t("install_libs") + "\n\n"
                    "python -m pip install python-docx reportlab arabic-reshaper python-bidi\n\n"
                    f"({t('details')}: {e})"
                )
            if saved_path:
                st.success(t("saved_ok", name=saved_path.name))

        others = {k: v for k, v in summaries.items() if k != sum_key}
        if others:
            open_set = st.session_state.setdefault("open_saved", set())
            with st.expander(t("others_exp", n=len(others))):
                for k, v in others.items():
                    is_open = k in open_set
                    icon = ":material/expand_less:" if is_open else ":material/expand_more:"
                    wide_button(
                        f"{icon}  {disp(v['table'])}",
                        key=f"open_{k}",
                        on_click=toggle_saved,
                        args=(k,),
                    )
                    if is_open:
                        st.caption(t("other_from", table=disp(v["table"]), file=v["file"]))
                        st.markdown(v["text"])
                        st.button(
                            "🗑", key=f"delsum_{k}", help=t("del_summary"),
                            on_click=delete_summary, args=(k,),
                        )
                        st.divider()

    # ---------- التبويب 5: الدردشة مع البيانات ----------
    if page == "chat":
        st.markdown(CHAT_CSS, unsafe_allow_html=True)
        st.subheader(t("chat_title"))
        st.caption(t("chat_hint"))

        if "messages" not in st.session_state:
            st.session_state.messages = []

        if st.button(t("clear_chat")):
            st.session_state.messages = []

        # حاوية للمحادثة، وخانة السؤال بتيجي بعدها تحت
        chat_box = st.container()
        prompt = st.chat_input(t("chat_input"), accept_audio=True)
        question, audio = None, None
        if prompt:
            if isinstance(prompt, str):
                question = prompt
            elif hasattr(prompt, "get"):
                question, audio = prompt.get("text"), prompt.get("audio")
            else:
                question = getattr(prompt, "text", None)
                audio = getattr(prompt, "audio", None)

        with chat_box:
            # عرض المحادثة السابقة (كل رسالة جنبها زرار نسخ وحذف بيظهروا لما تقفي عليها)
            for m in st.session_state.messages:
                render_message(m)

            if audio is not None and not question:
                if not api_key:
                    st.warning(t("need_key_short"))
                else:
                    try:
                        with st.spinner(t("spin_listen")):
                            question = transcribe(genai.Client(api_key=api_key), audio)
                    except Exception as e:
                        st.error(friendly_error(e))

            if question and not api_key:
                st.warning(t("need_key_short"))
            elif question:
                st.session_state.messages.append(
                    {"id": uuid.uuid4().hex, "role": "user", "content": question}
                )
                with st.chat_message("user"):
                    st.markdown(question)

                with st.chat_message("assistant"):
                    try:
                        client = genai.Client(api_key=api_key)
                        code_prompt = t(
                            "code_prompt",
                            dtypes=df.dtypes.to_string(),
                            head=df.head(5).to_string(),
                            question=question,
                        )

                        result, error = None, None
                        with st.spinner(t("spin_think")):
                            for attempt in range(2):
                                code = clean_code(ask_gemini(client, code_prompt))
                                try:
                                    result = run_code(code, df)
                                    error = None
                                    break
                                except Exception as e:
                                    error = str(e)
                                    code_prompt += t("fix_code", error=error)

                        if error:
                            answer = t("calc_fail")
                        else:
                            answer_prompt = t(
                                "answer_prompt", question=question, result=result_to_text(result)
                            )
                            try:
                                answer = ask_gemini(client, answer_prompt)
                            except Exception:
                                # الحصة خلصت؟ نعرض النتيجة المحسوبة من غير صياغة
                                answer = t("result_label") + "\n\n```\n" + result_to_text(result) + "\n```"

                        st.session_state.messages.append(
                            {"id": uuid.uuid4().hex, "role": "assistant", "content": answer}
                        )
                        st.rerun()  # نعيد الرسم عشان أزرار النسخ والحذف تظهر
                    except Exception as e:
                        if type(e).__name__ in ("RerunException", "StopException"):
                            raise
                        st.error(friendly_error(e))
